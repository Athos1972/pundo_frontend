#!/usr/bin/env python3
"""Build the three demo files of video #12 (catalog-import) from catalog-data.json.

    ~/dev/github/pundo_main_backend/.venv/bin/python video/catalog-import/make_catalog_files.py

niko-catalog.xlsx         24 rows, exactly one error: row 14 price "2,5O" (letter O instead of zero)
niko-catalog-fixed.xlsx   same list, row 14 corrected to 2.50
niko-catalog-update.xlsx  fixed list with one price changed and one row removed (scene "update later")
Columns and order = the template of /shop-admin/import (parsing.TEMPLATE_COLUMNS). Valid prices are numbers,
the faulty one is text, as it would be after a typo in Excel.
"""
import json
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent
DATA = json.loads((HERE / "catalog-data.json").read_text())


def build(name: str, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(DATA["columns"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        price = r["price"]
        try:
            price_cell: object = float(Decimal(price))
        except Exception:
            price_cell = price  # the typo stays text
        ws.append([r["sku"], r["name"], "en", r["ean"] or None, r["category"] or None, "true", price_cell, "EUR", r.get("image_url") or None])
    for col, width in zip("ABCDEFGHI", (12, 42, 9, 16, 34, 9, 8, 9, 60)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "0.00"
    wb.save(HERE / name)
    print(f"{name}: {len(rows)} rows")


rows = [dict(r) for r in DATA["rows"]]
err, upd = DATA["error"], DATA["update"]
broken = [dict(r, price=err["bad"]) if r["sku"] == err["sku"] else r for r in rows]
fixed = [dict(r, price=err["fixed"]) if r["sku"] == err["sku"] else r for r in rows]
update = [dict(r, price=upd["price"]["to"]) if r["sku"] == upd["price"]["sku"] else r
          for r in fixed if r["sku"] != upd["remove"]]
assert [r["sku"] for r in broken].index(err["sku"]) + 2 == 14, "error must sit in Excel row 14"
build("niko-catalog.xlsx", broken)
build("niko-catalog-fixed.xlsx", fixed)
build("niko-catalog-update.xlsx", update)
